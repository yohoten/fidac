#!/usr/bin/env python3
"""
17家A股车企 三大报表标准化采集
===============================
- 主研究对象: 赛力斯(601127)、长安汽车(000625)
- 对比对象: 比亚迪、长城、江淮、上汽、广汽、北汽蓝谷 + 9家 (共17家)
- 时间范围: 2016Q1 — 2025Q4 (含季度)
- 会计准则: 中国会计准则 (A股统一)
- 输出路径: FIDAC/data/
- 格式: 行=会计科目 | 列=报告期 | 值=元
"""

import sys
sys.stdout.reconfigure(encoding='utf-8')

import os, json, time, hashlib, logging
from pathlib import Path
from datetime import datetime

import pandas as pd, numpy as np
import akshare as ak
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

logging.basicConfig(level=logging.INFO, format='%(asctime)s [%(levelname)s] %(message)s', datefmt='%H:%M:%S')
logger = logging.getLogger(__name__)

# ======================== 路径 & 时间配置 ========================

BASE_DIR = Path(r'F:\（8）Desktop\财务数智决策应用赛_260725\FIDAC')
DATA_DIR = BASE_DIR / 'data'
CACHE_DIR = DATA_DIR / '.cache'
DATA_DIR.mkdir(parents=True, exist_ok=True)
CACHE_DIR.mkdir(parents=True, exist_ok=True)

TIME_START = 2016   # 起始年
TIME_END = 2025     # 结束年

# ======================== 17家A股车企 ========================

COMPANIES = [
    # 主研究对象
    ('601127', '赛力斯', '★主研究·AITO问界'),
    ('000625', '长安汽车', '★主研究·合资新能源'),

    # 乘用车龙头
    ('002594', '比亚迪', '纯电+插混'),
    ('600104', '上汽集团', '智己/荣威'),
    ('601633', '长城汽车', '哈弗/坦克'),
    ('601238', '广汽集团', '广汽埃安'),

    # 新能源乘用车
    ('600418', '江淮汽车', '蔚来代工·尊界'),
    ('600733', '北汽蓝谷', '极狐'),

    # 商用车 & 客车
    ('000550', '江铃汽车', '轻卡/皮卡'),
    ('000800', '一汽解放', '重卡龙头'),
    ('600006', '东风汽车', '商用车'),
    ('000951', '中国重汽', '重卡'),

    # 客车
    ('600066', '宇通客车', '客车龙头'),
    ('000957', '中通客车', '客车'),
    ('600686', '金龙汽车', '客车'),

    # 其他
    ('000572', '海马汽车', '乘用车'),
    ('601777', '力帆科技', '转型新能源'),
]

# ======================== 科目映射 ========================

ITEM_NAME_MAP = {
    'operating_income_total': '一、营业总收入',
    'operating_income': '营业收入',
    'operating_costs': '营业成本',
    'operating_costs_total': '营业成本',
    'taxes_and_surcharges': '税金及附加',
    'sales_fee': '销售费用',
    'manage_fee': '管理费用',
    'benefit_finance_fee': '财务费用',
    'financial_interest_expenses': '其中：利息费用',
    'financial_interest_income': '其中：利息收入',
    'research_and_development_expenses': '研发费用',
    'operating_profit': '营业利润',
    'other_income': '其他收益',
    'invest_income': '投资收益',
    'associate_invest_income': '其中：对联营企业和合营企业的投资收益',
    'fair_changes_income': '公允价值变动收益',
    'asset_disposal_income': '资产处置收益',
    'assets_impairment_loss': '资产减值损失',
    'benefit_credit_impairment_loss': '信用减值损失',
    'common_profit_total': '利润总额',
    'profit_total': '利润总额',
    'income_tax_expense': '所得税费用',
    'net_profit': '净利润',
    'continuing_net_profit': '持续经营净利润',
    'discontinued_operating_net_profit': '终止经营净利润',
    'parent_common_profit_total': '归属于母公司所有者的净利润',
    'parent_holder_net_profit': '归属于母公司所有者的净利润',
    'minority_holder_income_loss': '少数股东损益',
    'basic_eps': '基本每股收益', 'diluted_eps': '稀释每股收益',
    'non_operating_income': '营业外收入',
    'non_operating_expenses': '营业外支出',
    'other_common_profit': '其他综合收益',
    'minority_common_profit_total': '归属于少数股东的综合收益总额',
}

EXCLUDE_FROM_YUAN = {'基本每股收益', '稀释每股收益'}

# ======================== 值转换 ========================

def to_yuan(val):
    if val is None or val is np.nan or val is False:
        return None
    s = str(val).strip()
    if s.lower() in ('false', 'none', 'nan', 'n/a', '--', 'null', ''):
        return None
    try:
        s = s.replace(',', '').replace('−', '-')
        if '亿' in s: return float(s.replace('亿', '')) * 1e8
        elif '万' in s: return float(s.replace('万', '')) * 1e4
        elif '%' in s: return float(s.replace('%', ''))
        else: return float(s)
    except: return None

def clean_name(name):
    return str(name).strip().lstrip('*')

def period_label(date_str):
    try:
        dt = pd.to_datetime(date_str)
        y, m = dt.year, dt.month
        if m == 12: return f'{y}An'
        elif m == 9: return f'{y}Q3'
        elif m == 6: return f'{y}Q2'
        elif m == 3: return f'{y}Q1'
        else: return f'{y}M{m:02d}'
    except: return str(date_str)

def period_sort_key(label): return label
def period_display(label): return label.replace('An', '年报') if label else label

# ======================== 缓存 + 重试 ========================

def _cache_key(func_name, **kwargs):
    return hashlib.md5(f'{func_name}_{json.dumps(kwargs, sort_keys=True)}'.encode()).hexdigest()[:12]

def _load_cache(key):
    p = CACHE_DIR / f'{key}.pkl'
    return pd.read_pickle(p) if p.exists() else None

def _save_cache(key, df):
    if df is not None and len(df) > 0:
        try: df.to_pickle(CACHE_DIR / f'{key}.pkl')
        except: pass

def akshare_with_retry(func, max_retries=3, delay=1.0, **kwargs):
    key = _cache_key(func.__name__, **kwargs)
    cached = _load_cache(key)
    if cached is not None: return cached
    for attempt in range(max_retries):
        try:
            result = func(**kwargs)
            if result is not None and len(result) > 0:
                _save_cache(key, result)
                return result
            return result
        except Exception as e:
            if attempt < max_retries - 1:
                logger.warning(f'  R{attempt+1}: {e}')
                time.sleep(delay)
            else: logger.error(f'  FAIL: {e}')
            return None

# ======================== 三大报表获取 ========================

def _pivot_items(df, date_col, exclude_cols):
    items = {}
    df = df.copy()
    df['_date'] = pd.to_datetime(df[date_col], errors='coerce')
    df = df[df['_date'].notna()].sort_values('_date', ascending=True)
    # 过滤时间范围
    df = df[(df['_date'].dt.year >= TIME_START) & (df['_date'].dt.year <= TIME_END)]
    item_cols = [c for c in df.columns if c not in exclude_cols and c != '_date']
    for col in item_cols:
        vals = {}
        for _, row in df.iterrows():
            label = period_label(str(row['_date'].date()))
            v = to_yuan(row.get(col, None))
            if v is not None: vals[label] = v
        if vals: items[clean_name(col)] = vals
    return items


def fetch_and_pivot(symbol, name):
    logger.info(f'{name} ({symbol})')
    time.sleep(0.15)

    # 现金流量表
    df_cf = akshare_with_retry(ak.stock_financial_cash_ths, symbol=symbol, indicator='按报告期')
    cf_items = {}
    if df_cf is not None and len(df_cf) > 0:
        exclude = {df_cf.columns[0], '报告期', '报表核心指标'}
        cf_items = _pivot_items(df_cf, '报告期', exclude)
        logger.info(f'  CF: {len(cf_items)} 科目')
    time.sleep(0.1)

    # 资产负债表
    df_bs = akshare_with_retry(ak.stock_financial_debt_ths, symbol=symbol, indicator='按报告期')
    bs_items = {}
    if df_bs is not None and len(df_bs) > 0:
        exclude = {df_bs.columns[0], '报告期', '报表核心指标'}
        bs_items = _pivot_items(df_bs, '报告期', exclude)
        logger.info(f'  BS: {len(bs_items)} 科目')
    time.sleep(0.1)

    # 利润表
    df_is = akshare_with_retry(ak.stock_financial_benefit_new_ths, symbol=symbol, indicator='按报告期')
    is_items = {}
    if df_is is not None and len(df_is) > 0:
        df_is['_date'] = pd.to_datetime(df_is['report_date'], errors='coerce')
        df_is = df_is[df_is['_date'].notna()].sort_values('_date', ascending=True)
        # 过滤时间范围
        df_is = df_is[(df_is['_date'].dt.year >= TIME_START) & (df_is['_date'].dt.year <= TIME_END)]
        for metric, grp in df_is.groupby('metric_name'):
            cn_name = ITEM_NAME_MAP.get(metric, metric)
            vals = {}
            for _, row in grp.iterrows():
                label = period_label(str(row['_date'].date()))
                v = to_yuan(row.get('value', None))
                if v is not None: vals[label] = v
            if vals: is_items[cn_name] = vals
        logger.info(f'  IS: {len(is_items)} 科目')
    time.sleep(0.1)

    return cf_items, bs_items, is_items


# ======================== DataFrame 构建 ========================

def items_to_df(items_dict, row_order=None):
    if not items_dict: return pd.DataFrame()
    all_labels = set()
    for vals in items_dict.values(): all_labels.update(vals.keys())
    sorted_labels = sorted(all_labels, key=period_sort_key)
    if row_order:
        ordered = [n for n in row_order if n in items_dict]
        remaining = [n for n in items_dict if n not in ordered]
        row_names = ordered + remaining
    else: row_names = list(items_dict.keys())
    data = {label: [items_dict.get(name, {}).get(label, None) for name in row_names] for label in sorted_labels}
    df = pd.DataFrame(data, index=row_names)
    df.columns = [period_display(c) for c in df.columns]
    return df


# ======================== 标准科目顺序 ========================

CF_ORDER = [
    '一、经营活动产生的现金流量', '销售商品、提供劳务收到的现金',
    '收到的税费与返还', '收到其他与经营活动有关的现金', '经营活动现金流入小计',
    '购买商品、接受劳务支付的现金', '支付给职工以及为职工支付的现金',
    '支付的各项税费', '支付其他与经营活动有关的现金', '经营活动现金流出小计',
    '经营活动产生的现金流量净额',
    '二、投资活动产生的现金流量', '收回投资收到的现金', '取得投资收益收到的现金',
    '处置固定资产、无形资产和其他长期资产收回的现金净额',
    '收到其他与投资活动有关的现金', '投资活动现金流入小计',
    '购建固定资产、无形资产和其他长期资产支付的现金', '投资支付的现金',
    '投资活动现金流出小计', '投资活动产生的现金流量净额',
    '三、筹资活动产生的现金流量', '吸收投资收到的现金', '取得借款收到的现金',
    '筹资活动现金流入小计', '偿还债务支付的现金',
    '分配股利、利润或偿付利息支付的现金', '筹资活动现金流出小计',
    '筹资活动产生的现金流量净额',
    '现金及现金等价物净增加额', '加：期初现金及现金等价物余额',
    '期末现金及现金等价物余额',
]

BS_ORDER = [
    '一、流动资产', '货币资金', '交易性金融资产', '应收票据', '应收账款',
    '应收票据及应收账款', '预付款项', '其他应收款', '存货', '其他流动资产', '流动资产合计',
    '二、非流动资产', '固定资产', '在建工程', '无形资产', '长期股权投资',
    '商誉', '递延所得税资产', '其他非流动资产', '非流动资产合计',
    '三、资产合计',
    '四、流动负债', '短期借款', '应付票据', '应付账款', '预收款项',
    '合同负债', '应付职工薪酬', '应交税费', '其他应付款',
    '一年内到期的非流动负债', '其他流动负债', '流动负债合计',
    '五、非流动负债', '长期借款', '应付债券', '递延所得税负债',
    '其他非流动负债', '非流动负债合计',
    '六、负债合计',
    '七、所有者权益', '股本', '资本公积', '盈余公积', '未分配利润',
    '归属于母公司所有者权益合计', '少数股东权益', '所有者权益（或股东权益）合计',
]

IS_ORDER = [
    '一、营业总收入', '营业收入',
    '减：营业成本', '税金及附加', '销售费用', '管理费用', '研发费用',
    '财务费用', '其中：利息费用', '其中：利息收入',
    '加：其他收益', '投资收益', '其中：对联营企业和合营企业的投资收益',
    '公允价值变动收益', '资产处置收益', '资产减值损失', '信用减值损失',
    '营业利润', '加：营业外收入', '减：营业外支出',
    '利润总额', '减：所得税费用',
    '净利润', '持续经营净利润', '终止经营净利润',
    '归属于母公司所有者的净利润', '少数股东损益',
    '基本每股收益', '稀释每股收益',
]


# ======================== 辅助函数 ========================

def get_val(df, item_name, col):
    if item_name in df.index and col in df.columns:
        v = df.loc[item_name, col]
        if pd.notna(v) and v is not None: return float(v)
    return None

def safe_div(a, b):
    return a / b if a and b and b != 0 else None


# ======================== 派生指标（修复版）========================

def calc_derived_metrics(df_bs, df_is, df_cf):
    annual_cols = [c for c in df_is.columns if '年报' in str(c)]
    if not annual_cols: return pd.DataFrame()
    result = {}

    def set_metric(name, col, val):
        if name not in result: result[name] = {}
        result[name][col] = val

    for col in annual_cols:
        revenue = get_val(df_is, '营业收入', col)
        cost = get_val(df_is, '营业成本', col)
        net_profit = get_val(df_is, '净利润', col)
        profit_total = get_val(df_is, '利润总额', col)
        op_profit = get_val(df_is, '营业利润', col)
        rd = get_val(df_is, '研发费用', col)
        sales = get_val(df_is, '销售费用', col)
        admin = get_val(df_is, '管理费用', col)
        fin = get_val(df_is, '财务费用', col)
        parent_net = get_val(df_is, '归属于母公司所有者的净利润', col)

        total_assets = (get_val(df_bs, '资产合计', col) or get_val(df_bs, '*资产合计', col))
        total_liab = (get_val(df_bs, '负债合计', col) or get_val(df_bs, '*负债合计', col))
        equity = (get_val(df_bs, '所有者权益（或股东权益）合计', col)
                  or get_val(df_bs, '*所有者权益（或股东权益）合计', col))
        current_assets = get_val(df_bs, '流动资产合计', col)
        current_liab = get_val(df_bs, '流动负债合计', col)
        cash_eq = get_val(df_bs, '货币资金', col)
        receivables = (get_val(df_bs, '应收票据及应收账款', col) or get_val(df_bs, '应收账款', col))
        inventory = get_val(df_bs, '存货', col)
        fixed_assets = get_val(df_bs, '固定资产', col)

        prev_col = annual_cols[annual_cols.index(col) - 1] if annual_cols.index(col) > 0 else None
        prev_revenue = get_val(df_is, '营业收入', prev_col) if prev_col else None
        prev_assets = (get_val(df_bs, '资产合计', prev_col) or get_val(df_bs, '*资产合计', prev_col)) if prev_col else None
        prev_equity = (get_val(df_bs, '所有者权益（或股东权益）合计', prev_col)
                       or get_val(df_bs, '*所有者权益（或股东权益）合计', prev_col)) if prev_col else None
        prev_receivables = (get_val(df_bs, '应收票据及应收账款', prev_col)
                            or get_val(df_bs, '应收账款', prev_col)) if prev_col else None
        prev_inventory = get_val(df_bs, '存货', prev_col) if prev_col else None
        prev_current_assets = get_val(df_bs, '流动资产合计', prev_col) if prev_col else None
        prev_net = get_val(df_is, '净利润', prev_col) if prev_col else None

        op_cf = get_val(df_cf, '经营活动产生的现金流量净额', col)
        invest_cf = get_val(df_cf, '投资活动产生的现金流量净额', col)
        finance_cf = get_val(df_cf, '筹资活动产生的现金流量净额', col)
        capex = get_val(df_cf, '购建固定资产、无形资产和其他长期资产支付的现金', col)
        cash_from_sales = get_val(df_cf, '销售商品、提供劳务收到的现金', col)

        # 盈利能力
        gross_profit = revenue - cost if revenue and cost else None
        set_metric('毛利率(%)', col, safe_div(gross_profit, revenue) * 100 if gross_profit and revenue else None)
        set_metric('净利率(%)', col, safe_div(net_profit, revenue) * 100 if net_profit and revenue else None)
        set_metric('营业利润率(%)', col, safe_div(op_profit, revenue) * 100 if op_profit and revenue else None)

        avg_equity = (equity + prev_equity) / 2 if equity and prev_equity else equity
        avg_assets = (total_assets + prev_assets) / 2 if total_assets and prev_assets else total_assets
        set_metric('ROE(%)', col, safe_div(net_profit, avg_equity) * 100 if net_profit and avg_equity else None)
        set_metric('ROA(%)', col, safe_div(net_profit, avg_assets) * 100 if net_profit and avg_assets else None)
        set_metric('销售费用率(%)', col, safe_div(sales, revenue) * 100 if sales and revenue else None)
        set_metric('管理费用率(%)', col, safe_div(admin, revenue) * 100 if admin and revenue else None)
        set_metric('研发费用率(%)', col, safe_div(rd, revenue) * 100 if rd and revenue else None)
        set_metric('财务费用率(%)', col, safe_div(fin, revenue) * 100 if fin and revenue else None)

        # 偿债能力
        set_metric('资产负债率(%)', col, safe_div(total_liab, total_assets) * 100 if total_liab and total_assets else None)
        set_metric('流动比率', col, safe_div(current_assets, current_liab) if current_assets and current_liab else None)
        set_metric('速动比率', col, safe_div(current_assets - inventory, current_liab) if current_assets and inventory and current_liab else None)

        # 营运能力
        avg_receivables = (receivables + prev_receivables) / 2 if receivables and prev_receivables else receivables
        avg_inventory = (inventory + prev_inventory) / 2 if inventory and prev_inventory else inventory
        avg_ca = (current_assets + prev_current_assets) / 2 if current_assets and prev_current_assets else current_assets
        set_metric('应收账款周转率(次)', col, safe_div(revenue, avg_receivables) if revenue and avg_receivables else None)
        set_metric('应收账款周转天数(天)', col, 360 / safe_div(revenue, avg_receivables) if revenue and avg_receivables and safe_div(revenue, avg_receivables) else None)
        set_metric('存货周转率(次)', col, safe_div(cost, avg_inventory) if cost and avg_inventory else None)
        set_metric('存货周转天数(天)', col, 360 / safe_div(cost, avg_inventory) if cost and avg_inventory and safe_div(cost, avg_inventory) else None)
        set_metric('总资产周转率(次)', col, safe_div(revenue, avg_assets) if revenue and avg_assets else None)
        set_metric('流动资产周转率(次)', col, safe_div(revenue, avg_ca) if revenue and avg_ca else None)
        set_metric('固定资产周转率(次)', col, safe_div(revenue, fixed_assets) if revenue and fixed_assets else None)

        # 发展能力
        set_metric('营收同比(%)', col, (revenue - prev_revenue) / abs(prev_revenue) * 100 if revenue and prev_revenue and prev_revenue != 0 else None)
        set_metric('净利润同比(%)', col, (net_profit - prev_net) / abs(prev_net) * 100 if net_profit and prev_net and prev_net != 0 else None)
        set_metric('总资产增长率(%)', col, (total_assets - prev_assets) / abs(prev_assets) * 100 if total_assets and prev_assets and prev_assets != 0 else None)

        # 现金流质量
        set_metric('经营现金流/净利润', col, safe_div(op_cf, net_profit) if op_cf and net_profit else None)
        free_cf = op_cf - capex if op_cf and capex else None
        set_metric('自由现金流(元)', col, free_cf)
        set_metric('现金收入比', col, safe_div(cash_from_sales, revenue) if cash_from_sales and revenue else None)

        # 现金流组合
        combo = []
        for v in [op_cf, invest_cf, finance_cf]:
            if v is not None and v > 0: combo.append('+')
            elif v is not None and v < 0: combo.append('−')
            else: combo.append('?')
        set_metric('现金流组合', col, '/'.join(combo))

        # Altman Z-score
        x1 = safe_div(current_assets - current_liab, total_assets) if current_assets and current_liab and total_assets else None
        surplus = get_val(df_bs, '盈余公积', col) or 0
        retained = get_val(df_bs, '未分配利润', col) or 0
        x2 = safe_div(surplus + retained, total_assets) if total_assets else None
        interest_exp = get_val(df_is, '其中：利息费用', col) or fin or 0
        x3 = safe_div(profit_total + abs(interest_exp), total_assets) if profit_total and total_assets else None
        x4 = safe_div(equity, total_liab) if equity and total_liab else None
        x5 = safe_div(revenue, total_assets) if revenue and total_assets else None
        if all([x1, x2, x3, x4, x5]):
            z = 1.2 * x1 + 1.4 * x2 + 3.3 * x3 + 0.6 * x4 + 1.0 * x5
            set_metric('Altman Z-score', col, z)
            set_metric('Z-score判定', col, '安全区' if z > 2.99 else ('灰色区' if z >= 1.81 else '危险区'))

    return pd.DataFrame(result)


# ======================== Excel 写入 ========================

def write_statement_sheet(df, sheet_name, writer):
    if df is None or df.empty: return
    df.to_excel(writer, sheet_name=sheet_name[:31], startrow=0)
    ws = writer.sheets[sheet_name[:31]]

    hf = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    hfont = Font(bold=True, size=10, color='FFFFFF')
    sf = PatternFill(start_color='D9E2F3', end_color='D9E2F3', fill_type='solid')
    sfont = Font(bold=True, size=9, color='1F3864')
    wfont = Font(bold=True, size=9, color='C00000')
    b = Border(left=Side('thin','D9D9D9'), right=Side('thin','D9D9D9'),
               top=Side('thin','D9D9D9'), bottom=Side('thin','D9D9D9'))

    for col in range(1, ws.max_column + 1):
        c = ws.cell(row=1, column=col)
        c.font, c.fill, c.alignment = hfont, hf, Alignment(horizontal='center', vertical='center')
        c.border = b

    for row in range(2, ws.max_row + 1):
        ca = ws.cell(row=row, column=1)
        val = str(ca.value) if ca.value else ''
        if val.startswith(('一、','二、','三、','四、','五、','六、','七、')):
            ca.font, ca.fill = sfont, sf
        elif any(k in val for k in ['危险','预警']):
            ca.font = wfont
        else: ca.font = Font(size=9)
        ca.alignment = Alignment(horizontal='left', vertical='center')
        ca.border = b
        for col in range(2, ws.max_column + 1):
            c = ws.cell(row=row, column=col)
            c.alignment = Alignment(horizontal='right', vertical='center')
            c.border = b
            if c.value is not None and isinstance(c.value, (int, float)):
                rn = str(ws.cell(row=row, column=1).value) or ''
                if rn in EXCLUDE_FROM_YUAN: c.number_format = '0.0000'
                elif any(k in rn for k in ['率','比']): c.number_format = '0.00%'
                elif any(k in rn for k in ['次','天']): c.number_format = '0.00'
                elif abs(c.value) >= 1: c.number_format = '#,##0'
                else: c.number_format = '0.0000'
    ws.column_dimensions['A'].width = 50
    for col in range(2, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = 16
    ws.freeze_panes = 'B2'


# ======================== 单公司处理 ========================

def process_company(symbol, name):
    cf_items, bs_items, is_items = fetch_and_pivot(symbol, name)
    if not cf_items and not bs_items and not is_items:
        logger.warning(f'  ⚠️ 无数据')
        return None

    df_cf = items_to_df(cf_items, CF_ORDER)
    df_bs = items_to_df(bs_items, BS_ORDER)
    df_is = items_to_df(is_items, IS_ORDER)

    logger.info(f'  → CF:{df_cf.shape} BS:{df_bs.shape} IS:{df_is.shape}')

    df_derived = calc_derived_metrics(df_bs, df_is, df_cf)

    issues = []
    annual_cols = [c for c in df_bs.columns if '年报' in str(c)]
    for col in annual_cols:
        ta = get_val(df_bs, '资产合计', col) or get_val(df_bs, '*资产合计', col)
        tl = get_val(df_bs, '负债合计', col) or get_val(df_bs, '*负债合计', col)
        eq = get_val(df_bs, '所有者权益（或股东权益）合计', col) or get_val(df_bs, '*所有者权益（或股东权益）合计', col)
        if ta and tl and eq and abs(ta - (tl + eq)) / ta > 0.01:
            issues.append(f'[{col}] 差={abs(ta-tl-eq)/1e8:.1f}亿')

    output_path = DATA_DIR / f'{symbol}_{name}_财务数据.xlsx'
    writer = pd.ExcelWriter(output_path, engine='openpyxl')
    if not df_bs.empty: write_statement_sheet(df_bs, '资产负债表', writer)
    if not df_is.empty: write_statement_sheet(df_is, '利润表', writer)
    if not df_cf.empty: write_statement_sheet(df_cf, '现金流量表', writer)
    if not df_derived.empty: write_statement_sheet(df_derived, '财务分析指标', writer)
    df_issues = pd.DataFrame({'校验问题': issues}) if issues else pd.DataFrame({'校验结果': ['✅ 通过']})
    df_issues.to_excel(writer, sheet_name='数据质量校验', index=False)
    writer.close()
    logger.info(f'  ✅ {output_path.name} ({os.path.getsize(output_path)/1024:.0f}KB)')
    return output_path, issues


# ======================== 主流程 ========================

def main():
    logger.info(f'{"="*80}')
    logger.info(f'  17家A股车企 三大报表采集 (2016Q1—2025Q4)')
    logger.info(f'  会计准则: 中国会计准则 | 格式: 科目×日期 | 值=元')
    logger.info(f'  输出路径: {DATA_DIR}')
    logger.info(f'{"="*80}\n')

    results, all_issues = [], {}
    for symbol, name, note in COMPANIES:
        result = process_company(symbol, name)
        if result:
            path, issues = result
            results.append((symbol, name, note, path))
            all_issues[name] = issues

    # 综合对比表
    logger.info(f'\n{"="*80}')
    logger.info(f'  综合对比表 (17家车企)')
    logger.info(f'{"="*80}')

    comp_rows = []
    for symbol, name, note, path in results:
        xl = pd.ExcelFile(path)
        if '利润表' in xl.sheet_names:
            df_is = pd.read_excel(xl, '利润表', index_col=0)
            annual_cols = [c for c in df_is.columns if '年报' in str(c)]
            if annual_cols:
                latest = annual_cols[-1]
                row = {'代码': symbol, '公司': name, '分类': note, '最新年报': latest}
                for item in ['营业收入','净利润','归属于母公司所有者的净利润',
                            '营业利润','利润总额','基本每股收益','稀释每股收益',
                            '营业成本','销售费用','管理费用','研发费用','财务费用']:
                    if item in df_is.index:
                        v = df_is.loc[item, latest]
                        if pd.notna(v): row[item] = v
                if '财务分析指标' in xl.sheet_names:
                    df_d = pd.read_excel(xl, '财务分析指标', index_col=0)
                    for m in ['毛利率(%)','净利率(%)','ROE(%)','资产负债率(%)',
                              '经营现金流/净利润','Altman Z-score']:
                        if m in df_d.index and latest in df_d.columns:
                            v = df_d.loc[m, latest]
                            if pd.notna(v): row[m] = v
                comp_rows.append(row)

    if comp_rows:
        df_comp = pd.DataFrame(comp_rows)
        if '营业收入' in df_comp.columns:
            df_comp = df_comp.sort_values('营业收入', ascending=False).reset_index(drop=True)
        df_comp.to_excel(DATA_DIR / '17家车企综合对比.xlsx', index=False)
        logger.info(f'  ✅ 17家车企综合对比.xlsx')

        # 打印速览
        key_cols = ['代码','公司','最新年报','营业收入','净利润','毛利率(%)','净利率(%)','ROE(%)']
        avail = [c for c in key_cols if c in df_comp.columns]
        logger.info(f'\n{"="*80}')
        logger.info(f'  17家车企 最新年报速览 (按营收排序)')
        logger.info(f'{"="*80}')
        df_show = df_comp[avail].copy()
        for c in df_show.columns:
            if c in ('代码','公司','最新年报'): continue
            df_show[c] = df_show[c].apply(lambda x: f'{x/1e8:.1f}亿' if pd.notna(x) and abs(x)>1e6
                                          else (f'{x:.2f}%' if pd.notna(x) and '%' in c else (f'{x:.4f}' if pd.notna(x) else '-')))
        print(df_show.to_string(index=False))

    # 数据质量
    logger.info(f'\n{"="*80}')
    logger.info(f'  数据质量报告')
    logger.info(f'{"="*80}')
    for n, issues in all_issues.items():
        status = '✅' if not issues else f'⚠️ {len(issues)}'
        logger.info(f'  {n}: {status}')
        for i in issues[:1]: logger.info(f'    {i}')

    logger.info(f'\n全部文件: {DATA_DIR}')


if __name__ == '__main__':
    main()
