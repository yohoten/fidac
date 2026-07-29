#!/usr/bin/env python3
"""
2026Q1 MAPE 数据追加采集（增量模式）— 赛力斯 & 长安汽车
=====================================================
用途: 从 AkShare 获取 2026Q1 数据，增量追加到现有 Excel，
      保留已有「合并报表」「财务指标表」等所有 sheet。

使用方法:
  cd F:/FIDAC/data
  python collect_2026q1_for_mape.py

依赖: pip install akshare openpyxl
"""

import sys, os, time, hashlib, json, logging
from pathlib import Path

import pandas as pd
import numpy as np
import akshare as ak

# ── 路径 ──
DATA_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(DATA_DIR.parent))  # FIDAC 根目录
from lib.financial_metrics import integrate_financial_data, compute_ratios

logging.basicConfig(level=logging.INFO,
                    format='%(asctime)s [%(levelname)s] %(message)s',
                    datefmt='%H:%M:%S')
logger = logging.getLogger(__name__)

# ======================== 参数 ========================

TARGETS = [
    ('601127', '赛力斯'),
    ('000625', '长安汽车'),
]

METRICS_OF_INTEREST = ['毛利率', '资产负债率']

# ======================== AkShare 工具函数 ========================

def to_yuan(val):
    if val is None or val is np.nan or val is False:
        return None
    s = str(val).strip()
    if s.lower() in ('false', 'none', 'nan', 'n/a', '--', 'null', ''):
        return None
    try:
        s = s.replace(',', '').replace('−', '-')
        if '亿' in s: return float(s.replace('亿', '')) * 1e8
        if '万' in s: return float(s.replace('万', '')) * 1e4
        if '%' in s: return float(s.replace('%', ''))
        return float(s)
    except:
        return None

def period_label(d):
    try:
        dt = pd.to_datetime(d)
        y, m = dt.year, dt.month
        if m == 12: return f'{y}年报'
        if m == 9: return f'{y}Q3'
        if m == 6: return f'{y}Q2'
        if m == 3: return f'{y}Q1'
        return f'{y}M{m:02d}'
    except:
        return str(d)

def _cache_key(func_name, **kw):
    raw = f'{func_name}_{json.dumps(kw, sort_keys=True)}'
    return hashlib.md5(raw.encode()).hexdigest()[:12]

def _load_cache(k):
    cache_dir = DATA_DIR / '.cache'
    p = cache_dir / f'{k}.pkl'
    return pd.read_pickle(p) if p.exists() else None

def _save_cache(k, df):
    cache_dir = DATA_DIR / '.cache'
    cache_dir.mkdir(exist_ok=True)
    if df is not None and len(df) > 0:
        try:
            df.to_pickle(cache_dir / f'{k}.pkl')
        except:
            pass

def akshare_r(func, max_retries=3, delay=1.0, **kw):
    k = _cache_key(func.__name__, **kw)
    c = _load_cache(k)
    if c is not None:
        return c
    for a in range(max_retries):
        try:
            r = func(**kw)
            if r is not None and len(r) > 0:
                _save_cache(k, r)
                return r
            return r
        except Exception as e:
            if a < max_retries - 1:
                time.sleep(delay)
            else:
                logger.error(f'  FAIL({func.__name__}): {e}')
                return None

# ======================== 增量采集 2026Q1 ========================

def collect_2026q1_col(symbol, name):
    """从 AkShare 获取 2026Q1 列数据，返回 {sheet_name: {科目: 值}}"""
    logger.info(f'  [{name}] 增量采集 2026Q1 ...')

    result = {'资产负债表': {}, '利润表': {}, '现金流量表': {}}
    time.sleep(0.5)

    # ── 资产负债表 ──
    try:
        df_bs = akshare_r(ak.stock_financial_debt_ths, symbol=symbol, indicator='按报告期')
        if df_bs is not None and len(df_bs) > 0:
            df_bs['_date'] = pd.to_datetime(df_bs['报告期'], errors='coerce')
            mask = (df_bs['_date'].dt.year == 2026) & (df_bs['_date'].dt.month.isin([3]))
            df_q1 = df_bs[mask]
            exclude = {df_bs.columns[0], '报告期', '报表核心指标'}
            for col in [c for c in df_bs.columns if c not in exclude and c != '_date']:
                vals = df_q1[col].dropna()
                if len(vals) > 0:
                    v = to_yuan(vals.iloc[0])
                    if v is not None:
                        result['资产负债表'][col.strip().lstrip('*')] = v
            logger.info(f'    资产负债表: {len(result["资产负债表"])} 科目')
    except Exception as e:
        logger.warning(f'    资产负债表失败: {e}')

    time.sleep(0.3)

    # ── 利润表 ──
    try:
        df_is = akshare_r(ak.stock_financial_benefit_new_ths, symbol=symbol, indicator='按报告期')
        if df_is is not None and len(df_is) > 0:
            df_is['_date'] = pd.to_datetime(df_is['report_date'], errors='coerce')
            mask = (df_is['_date'].dt.year == 2026) & (df_is['_date'].dt.month.isin([3]))
            df_q1 = df_is[mask]
            for _, row in df_q1.iterrows():
                metric = row.get('metric_name', '')
                v = to_yuan(row.get('value'))
                if v is not None:
                    # 使用中文映射名
                    cn_map = {
                        'operating_income_total': '一、营业总收入',
                        'operating_income': '营业收入',
                        'operating_costs': '营业成本',
                        'operating_costs_total': '营业成本',
                        'taxes_and_surcharges': '税金及附加',
                        'sales_fee': '销售费用',
                        'manage_fee': '管理费用',
                        'benefit_finance_fee': '财务费用',
                        'financial_interest_expenses': '其中：利息费用',
                        'financial_interest_income': '其中：利息收入',
                        'research_and_development_expenses': '研发费用',
                        'operating_profit': '营业利润',
                        'other_income': '其他收益',
                        'invest_income': '投资收益',
                        'fair_changes_income': '公允价值变动收益',
                        'asset_disposal_income': '资产处置收益',
                        'assets_impairment_loss': '资产减值损失',
                        'benefit_credit_impairment_loss': '信用减值损失',
                        'common_profit_total': '利润总额',
                        'profit_total': '利润总额',
                        'income_tax_expense': '所得税费用',
                        'net_profit': '净利润',
                        'parent_common_profit_total': '归属于母公司所有者的净利润',
                        'parent_holder_net_profit': '归属于母公司所有者的净利润',
                        'minority_holder_income_loss': '少数股东损益',
                        'basic_eps': '基本每股收益',
                        'diluted_eps': '稀释每股收益',
                        'non_operating_income': '营业外收入',
                        'non_operating_expenses': '营业外支出',
                    }
                    cn = cn_map.get(metric, metric)
                    result['利润表'][cn] = v
            logger.info(f'    利润表: {len(result["利润表"])} 科目')
    except Exception as e:
        logger.warning(f'    利润表失败: {e}')

    time.sleep(0.3)

    # ── 现金流量表 ──
    try:
        df_cf = akshare_r(ak.stock_financial_cash_ths, symbol=symbol, indicator='按报告期')
        if df_cf is not None and len(df_cf) > 0:
            df_cf['_date'] = pd.to_datetime(df_cf['报告期'], errors='coerce')
            mask = (df_cf['_date'].dt.year == 2026) & (df_cf['_date'].dt.month.isin([3]))
            df_q1 = df_cf[mask]
            exclude = {df_cf.columns[0], '报告期', '报表核心指标'}
            for col in [c for c in df_cf.columns if c not in exclude and c != '_date']:
                vals = df_q1[col].dropna()
                if len(vals) > 0:
                    v = to_yuan(vals.iloc[0])
                    if v is not None:
                        result['现金流量表'][col.strip().lstrip('*')] = v
            logger.info(f'    现金流量表: {len(result["现金流量表"])} 科目')
    except Exception as e:
        logger.warning(f'    现金流量表失败: {e}')

    return result

# ======================== 追加写入 Excel ========================

def append_to_excel(filepath, code, name, new_data):
    """
    将 2026Q1 数据追加到现有 Excel，保留所有 sheet。
    同时重新生成「合并报表」和「财务指标表」。
    """
    from openpyxl import load_workbook

    wb = load_workbook(filepath)
    target_col = '2026Q1'

    # Step 1: 追加列到三大报表 sheet
    for sheet_name in ['资产负债表', '利润表', '现金流量表']:
        if sheet_name not in new_data or not new_data[sheet_name]:
            continue
        if sheet_name not in wb.sheetnames:
            continue

        ws = wb[sheet_name]
        # 找到最后一列的位置
        max_col = ws.max_column
        # 在最后一列后追加新列
        new_col_idx = max_col + 1
        ws.cell(row=1, column=new_col_idx, value=target_col)

        # 建立科目 → 行号的映射
        row_map = {}
        for row in range(2, ws.max_row + 1):
            cell_val = str(ws.cell(row=row, column=1).value or '').strip().lstrip('*')
            if cell_val:
                row_map[cell_val] = row

        count = 0
        for item, val in new_data[sheet_name].items():
            item_clean = item.strip().lstrip('*')
            if item_clean in row_map:
                ws.cell(row=row_map[item_clean], column=new_col_idx, value=val)
                count += 1
            else:
                # 不在映射中，追加为新行
                new_row = ws.max_row + 1
                ws.cell(row=new_row, column=1, value=item)
                ws.cell(row=new_row, column=new_col_idx, value=val)
                count += 1

        logger.info(f'    {sheet_name}: 写入 {count} 个科目到 {target_col}')

    wb.save(filepath)
    wb.close()

    # Step 2: 用 pandas 重新生成「合并报表」
    integrated = integrate_financial_data(code, data_dir=str(DATA_DIR))
    with pd.ExcelWriter(filepath, engine='openpyxl', mode='a', if_sheet_exists='replace') as writer:
        integrated.to_excel(writer, sheet_name='合并报表', index=True)
    logger.info(f'    合并报表: {len(integrated)} 期 × {len(integrated.columns)} 指标')

    # Step 3: 重新生成「财务指标表」
    ratios = compute_ratios(integrated)
    with pd.ExcelWriter(filepath, engine='openpyxl', mode='a', if_sheet_exists='replace') as writer:
        ratios.to_excel(writer, sheet_name='财务指标表', index=True)
    logger.info(f'    财务指标表: {len(ratios)} 期 × {len(ratios.columns)} 指标')

    # Step 4: 提取目标指标实际值
    actuals = {}
    if '2026Q1' in ratios.index:
        for m in METRICS_OF_INTEREST:
            if m in ratios.columns:
                v = ratios.loc['2026Q1', m]
                if pd.notna(v):
                    actuals[m] = round(float(v), 6)

    return actuals

# ======================== 主入口 ========================

def main():
    logger.info(f'\n{"="*60}')
    logger.info(f'  2026Q1 MAPE 数据追加采集（增量模式）')
    logger.info(f'  目标: 赛力斯(601127) & 长安汽车(000625)')
    logger.info(f'  模式: 只追加 2026Q1 列，保留现有所有 sheet')
    logger.info(f'{"="*60}')

    all_actuals = {}

    for symbol, name in TARGETS:
        code = f'{symbol}_{name}'
        filepath = DATA_DIR / f'{code}_财务数据.xlsx'

        if not filepath.exists():
            logger.error(f'  ⚠️ {filepath.name} 不存在，跳过')
            continue

        logger.info(f'\n  ── 处理: {name} ({symbol}) ──')

        # 增量采集
        new_data = collect_2026q1_col(symbol, name)

        if not any(new_data.values()):
            logger.warning(f'    ⚠️ {name}: 未获取到任何 2026Q1 数据（AkShare 可能尚未发布）')
            continue

        # 追加写入
        actuals = append_to_excel(filepath, code, name, new_data)
        if actuals:
            all_actuals[name] = actuals

    # ─── 摘要 ───
    logger.info(f'\n{"="*60}')
    logger.info(f'  📊 2026Q1 实际值摘要')
    logger.info(f'{"="*60}')

    if all_actuals:
        for name, vals in all_actuals.items():
            logger.info(f'  {name}:')
            for m, v in vals.items():
                logger.info(f'    {m} = {v:.6f} ({v*100:.2f}%)')
        logger.info(f'\n  📋 下一步: 直接在 notebook 重新运行 Prophet Cell + MAPE Cell 即可')
    else:
        logger.warning(f'  ⚠️ 未能提取任何 2026Q1 数据')
        logger.warning(f'  可能原因: AkShare 尚未发布 2026Q1 财报数据')

    return all_actuals


if __name__ == '__main__':
    main()
