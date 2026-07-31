#!/usr/bin/env python3
"""
财务大数据智能决策竞赛 — 统一数据采集脚本
==========================================
模块1: 8家A股乘用车车企 三大报表标准化采集 (2016Q1—2025Q4)
模块2: 中国汽车行业多维销量数据采集 (乘联会月度)
模块3: 综合对比表 & 数据质量报告
输出路径: FIDAC/data/
会计准则: 中国会计准则(A股统一) | 格式: 科目×日期 | 值=元
"""

import sys
sys.stdout.reconfigure(encoding='utf-8')

import os, json, time, hashlib, logging
from pathlib import Path

import pandas as pd, numpy as np
import akshare as ak
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

logging.basicConfig(level=logging.INFO, format='%(asctime)s [%(levelname)s] %(message)s', datefmt='%H:%M:%S')
logger = logging.getLogger(__name__)

# ======================== 路径 & 参数 ========================

BASE_DIR = Path(r'F:\（8）Desktop\财务数智决策应用赛_260725\FIDAC')
DATA_DIR = BASE_DIR / 'data'
CACHE_DIR = DATA_DIR / '.cache'
DATA_DIR.mkdir(parents=True, exist_ok=True)
CACHE_DIR.mkdir(parents=True, exist_ok=True)

TIME_START, TIME_END = 2016, 2025
TIME_END_EXTENDED = 2026  # 用于 MAPE 评估时扩展到 2026Q1

# ======================== 8家A股乘用车车企 ========================
# 聚焦乘用车整车企业, 统一会计准则, 2016Q1-2025Q4

COMPANIES = [
    # ★ 主研究对象 (重庆本土)
    ('601127','赛力斯','★主研究·AITO问界'),    ('000625','长安汽车','★主研究·深蓝/阿维塔'),

    # 乘用车龙头
    ('002594','比亚迪','纯电+插混龙头'),       ('600104','上汽集团','智己/荣威/MG'),
    ('601633','长城汽车','哈弗/坦克/魏牌'),     ('601238','广汽集团','广汽埃安/传祺'),

    # 新能源乘用车
    ('600418','江淮汽车','蔚来代工·尊界合作'),   ('600733','北汽蓝谷','极狐'),
]

# ======================== 科目映射 ========================

ITEM_NAME_MAP = {
    'operating_income_total':'一、营业总收入','operating_income':'营业收入',
    'operating_costs':'营业成本','operating_costs_total':'营业成本',
    'taxes_and_surcharges':'税金及附加','sales_fee':'销售费用',
    'manage_fee':'管理费用','benefit_finance_fee':'财务费用',
    'financial_interest_expenses':'其中：利息费用','financial_interest_income':'其中：利息收入',
    'research_and_development_expenses':'研发费用','operating_profit':'营业利润',
    'other_income':'其他收益','invest_income':'投资收益',
    'associate_invest_income':'其中：对联营企业和合营企业的投资收益',
    'fair_changes_income':'公允价值变动收益','asset_disposal_income':'资产处置收益',
    'assets_impairment_loss':'资产减值损失','benefit_credit_impairment_loss':'信用减值损失',
    'common_profit_total':'利润总额','profit_total':'利润总额',
    'income_tax_expense':'所得税费用','net_profit':'净利润',
    'continuing_net_profit':'持续经营净利润','discontinued_operating_net_profit':'终止经营净利润',
    'parent_common_profit_total':'归属于母公司所有者的净利润',
    'parent_holder_net_profit':'归属于母公司所有者的净利润',
    'minority_holder_income_loss':'少数股东损益',
    'basic_eps':'基本每股收益','diluted_eps':'稀释每股收益',
    'non_operating_income':'营业外收入','non_operating_expenses':'营业外支出',
    'other_common_profit':'其他综合收益',
}

EXCLUDE_FROM_YUAN = {'基本每股收益','稀释每股收益'}

# ======================== 工具函数 ========================

def to_yuan(val):
    if val is None or val is np.nan or val is False: return None
    s = str(val).strip()
    if s.lower() in ('false','none','nan','n/a','--','null',''): return None
    try:
        s = s.replace(',','').replace('−','-')
        if '亿' in s: return float(s.replace('亿','')) * 1e8
        if '万' in s: return float(s.replace('万','')) * 1e4
        if '%' in s: return float(s.replace('%',''))
        return float(s)
    except: return None

def clean_name(n): return str(n).strip().lstrip('*')

def period_label(d):
    try:
        dt = pd.to_datetime(d); y,m = dt.year, dt.month
        if m==12: return f'{y}An'
        if m==9: return f'{y}Q3'
        if m==6: return f'{y}Q2'
        if m==3: return f'{y}Q1'
        return f'{y}M{m:02d}'
    except: return str(d)

def period_display(l): return l.replace('An','年报') if l else l
def period_sort_key(l): return l
def safe_div(a,b): return a/b if a and b and b!=0 else None

# ======================== 缓存 ========================

def _cache_key(func_name,**kw):
    return hashlib.md5(f'{func_name}_{json.dumps(kw,sort_keys=True)}'.encode()).hexdigest()[:12]

def _load_cache(k):
    p = CACHE_DIR / f'{k}.pkl'
    return pd.read_pickle(p) if p.exists() else None

def _save_cache(k,df):
    if df is not None and len(df)>0:
        try: df.to_pickle(CACHE_DIR/f'{k}.pkl')
        except: pass

def akshare_r(func,max_retries=3,delay=1.0,**kw):
    k = _cache_key(func.__name__,**kw)
    c = _load_cache(k)
    if c is not None: return c
    for a in range(max_retries):
        try:
            r = func(**kw)
            if r is not None and len(r)>0: _save_cache(k,r); return r
            return r
        except Exception as e:
            if a<max_retries-1: time.sleep(delay)
            else: logger.error(f'  FAIL({func.__name__}): {e}'); return None

# ======================== 模块1: 三大报表 ========================

def _pivot_items(df, date_col, exclude_cols):
    items = {}
    df = df.copy()
    df['_date'] = pd.to_datetime(df[date_col], errors='coerce')
    df = df[df['_date'].notna()].sort_values('_date', ascending=True)
    df = df[(df['_date'].dt.year >= TIME_START) & (df['_date'].dt.year <= TIME_END)]
    for col in [c for c in df.columns if c not in exclude_cols and c!='_date']:
        vals = {}
        for _,row in df.iterrows():
            v = to_yuan(row.get(col))
            if v is not None: vals[period_label(str(row['_date'].date()))] = v
        if vals: items[clean_name(col)] = vals
    return items


def fetch_and_pivot(symbol, name):
    logger.info(f'  {name} ({symbol})')
    time.sleep(0.12)

    # CF
    df_cf = akshare_r(ak.stock_financial_cash_ths, symbol=symbol, indicator='按报告期')
    cf = {}
    if df_cf is not None and len(df_cf)>0:
        ex = {df_cf.columns[0],'报告期','报表核心指标'}
        cf = _pivot_items(df_cf,'报告期',ex)
    time.sleep(0.08)

    # BS
    df_bs = akshare_r(ak.stock_financial_debt_ths, symbol=symbol, indicator='按报告期')
    bs = {}
    if df_bs is not None and len(df_bs)>0:
        ex = {df_bs.columns[0],'报告期','报表核心指标'}
        bs = _pivot_items(df_bs,'报告期',ex)
    time.sleep(0.08)

    # IS
    df_is = akshare_r(ak.stock_financial_benefit_new_ths, symbol=symbol, indicator='按报告期')
    is_ = {}
    if df_is is not None and len(df_is)>0:
        df_is['_date'] = pd.to_datetime(df_is['report_date'], errors='coerce')
        df_is = df_is[df_is['_date'].notna()].sort_values('_date', ascending=True)
        df_is = df_is[(df_is['_date'].dt.year>=TIME_START)&(df_is['_date'].dt.year<=TIME_END)]
        for metric,grp in df_is.groupby('metric_name'):
            cn = ITEM_NAME_MAP.get(metric, metric)
            vals = {}
            for _,row in grp.iterrows():
                v = to_yuan(row.get('value'))
                if v is not None: vals[period_label(str(row['_date'].date()))] = v
            if vals: is_[cn] = vals
    time.sleep(0.08)
    return cf,bs,is_


def items_to_df(items_dict, row_order=None):
    if not items_dict: return pd.DataFrame()
    labels = set()
    for vs in items_dict.values(): labels.update(vs.keys())
    sorted_l = sorted(labels, key=period_sort_key)
    if row_order:
        ordered = [n for n in row_order if n in items_dict]
        remaining = [n for n in items_dict if n not in ordered]
        names = ordered + remaining
    else: names = list(items_dict.keys())
    data = {l: [items_dict.get(n,{}).get(l) for n in names] for l in sorted_l}
    df = pd.DataFrame(data, index=names)
    df.columns = [period_display(c) for c in df.columns]
    return df


def get_val(df, item, col):
    if item in df.index and col in df.columns:
        v = df.loc[item, col]
        if pd.notna(v) and v is not None: return float(v)
    return None

# ======================== 标准科目顺序 ========================

CF_ORDER = [
    '一、经营活动产生的现金流量','销售商品、提供劳务收到的现金',
    '收到的税费与返还','收到其他与经营活动有关的现金','经营活动现金流入小计',
    '购买商品、接受劳务支付的现金','支付给职工以及为职工支付的现金',
    '支付的各项税费','支付其他与经营活动有关的现金','经营活动现金流出小计',
    '经营活动产生的现金流量净额',
    '二、投资活动产生的现金流量','收回投资收到的现金','取得投资收益收到的现金',
    '处置固定资产、无形资产和其他长期资产收回的现金净额',
    '收到其他与投资活动有关的现金','投资活动现金流入小计',
    '购建固定资产、无形资产和其他长期资产支付的现金','投资支付的现金',
    '投资活动现金流出小计','投资活动产生的现金流量净额',
    '三、筹资活动产生的现金流量','吸收投资收到的现金','取得借款收到的现金',
    '筹资活动现金流入小计','偿还债务支付的现金',
    '分配股利、利润或偿付利息支付的现金','筹资活动现金流出小计',
    '筹资活动产生的现金流量净额',
    '现金及现金等价物净增加额','加：期初现金及现金等价物余额','期末现金及现金等价物余额',
]

BS_ORDER = [
    '一、流动资产','货币资金','交易性金融资产','应收票据','应收账款',
    '应收票据及应收账款','预付款项','其他应收款','存货','其他流动资产','流动资产合计',
    '二、非流动资产','固定资产','在建工程','无形资产','长期股权投资',
    '商誉','递延所得税资产','其他非流动资产','非流动资产合计',
    '三、资产合计',
    '四、流动负债','短期借款','应付票据','应付账款','预收款项',
    '合同负债','应付职工薪酬','应交税费','其他应付款',
    '一年内到期的非流动负债','其他流动负债','流动负债合计',
    '五、非流动负债','长期借款','应付债券','递延所得税负债',
    '其他非流动负债','非流动负债合计',
    '六、负债合计',
    '七、所有者权益','股本','资本公积','盈余公积','未分配利润',
    '归属于母公司所有者权益合计','少数股东权益','所有者权益（或股东权益）合计',
]

IS_ORDER = [
    '一、营业总收入','营业收入',
    '减：营业成本','税金及附加','销售费用','管理费用','研发费用',
    '财务费用','其中：利息费用','其中：利息收入',
    '加：其他收益','投资收益','其中：对联营企业和合营企业的投资收益',
    '公允价值变动收益','资产处置收益','资产减值损失','信用减值损失',
    '营业利润','加：营业外收入','减：营业外支出',
    '利润总额','减：所得税费用',
    '净利润','持续经营净利润','终止经营净利润',
    '归属于母公司所有者的净利润','少数股东损益',
    '基本每股收益','稀释每股收益',
]

# ======================== 派生指标 ========================

def calc_derived_metrics(df_bs, df_is, df_cf):
    annual_cols = [c for c in df_is.columns if '年报' in str(c)]
    if not annual_cols: return pd.DataFrame()
    result = {}

    def s(n,c,v):
        if n not in result: result[n] = {}
        result[n][c] = v

    for col in annual_cols:
        rev = get_val(df_is,'营业收入',col); cost = get_val(df_is,'营业成本',col)
        np_ = get_val(df_is,'净利润',col); pt = get_val(df_is,'利润总额',col)
        op = get_val(df_is,'营业利润',col)
        rd = get_val(df_is,'研发费用',col); sales = get_val(df_is,'销售费用',col)
        admin = get_val(df_is,'管理费用',col); fin = get_val(df_is,'财务费用',col)
        pn = get_val(df_is,'归属于母公司所有者的净利润',col)

        ta = get_val(df_bs,'资产合计',col) or get_val(df_bs,'*资产合计',col)
        tl = get_val(df_bs,'负债合计',col) or get_val(df_bs,'*负债合计',col)
        eq = get_val(df_bs,'所有者权益（或股东权益）合计',col) or get_val(df_bs,'*所有者权益（或股东权益）合计',col)
        ca = get_val(df_bs,'流动资产合计',col); cl = get_val(df_bs,'流动负债合计',col)
        recv = get_val(df_bs,'应收票据及应收账款',col) or get_val(df_bs,'应收账款',col)
        inv = get_val(df_bs,'存货',col); fa = get_val(df_bs,'固定资产',col)

        pc = annual_cols[annual_cols.index(col)-1] if annual_cols.index(col)>0 else None
        prev_rev = get_val(df_is,'营业收入',pc) if pc else None
        prev_ta = (get_val(df_bs,'资产合计',pc) or get_val(df_bs,'*资产合计',pc)) if pc else None
        prev_eq = (get_val(df_bs,'所有者权益（或股东权益）合计',pc) or get_val(df_bs,'*所有者权益（或股东权益）合计',pc)) if pc else None
        prev_recv = (get_val(df_bs,'应收票据及应收账款',pc) or get_val(df_bs,'应收账款',pc)) if pc else None
        prev_inv = get_val(df_bs,'存货',pc) if pc else None
        prev_ca = get_val(df_bs,'流动资产合计',pc) if pc else None
        prev_np = get_val(df_is,'净利润',pc) if pc else None

        op_cf = get_val(df_cf,'经营活动产生的现金流量净额',col)
        invest_cf = get_val(df_cf,'投资活动产生的现金流量净额',col)
        finance_cf = get_val(df_cf,'筹资活动产生的现金流量净额',col)
        capex = get_val(df_cf,'购建固定资产、无形资产和其他长期资产支付的现金',col)
        cfs = get_val(df_cf,'销售商品、提供劳务收到的现金',col)

        gp = rev-cost if rev and cost else None
        avg_eq = (eq+prev_eq)/2 if eq and prev_eq else eq
        avg_ta = (ta+prev_ta)/2 if ta and prev_ta else ta
        avg_recv = (recv+prev_recv)/2 if recv and prev_recv else recv
        avg_inv = (inv+prev_inv)/2 if inv and prev_inv else inv
        avg_ca = (ca+prev_ca)/2 if ca and prev_ca else ca

        # 盈利能力
        s('毛利率(%)',col,safe_div(gp,rev)*100 if gp and rev else None)
        s('净利率(%)',col,safe_div(np_,rev)*100 if np_ and rev else None)
        s('营业利润率(%)',col,safe_div(op,rev)*100 if op and rev else None)
        s('ROE(%)',col,safe_div(np_,avg_eq)*100 if np_ and avg_eq else None)
        s('ROA(%)',col,safe_div(np_,avg_ta)*100 if np_ and avg_ta else None)
        s('销售费用率(%)',col,safe_div(sales,rev)*100 if sales and rev else None)
        s('管理费用率(%)',col,safe_div(admin,rev)*100 if admin and rev else None)
        s('研发费用率(%)',col,safe_div(rd,rev)*100 if rd and rev else None)

        # 偿债
        s('资产负债率(%)',col,safe_div(tl,ta)*100 if tl and ta else None)
        s('流动比率',col,safe_div(ca,cl) if ca and cl else None)
        s('速动比率',col,safe_div(ca-inv,cl) if ca and inv and cl else None)
        s('产权比率(%)',col,safe_div(tl,eq)*100 if tl and eq else None)

        # 营运
        s('应收账款周转率(次)',col,safe_div(rev,avg_recv) if rev and avg_recv else None)
        s('存货周转率(次)',col,safe_div(cost,avg_inv) if cost and avg_inv else None)
        s('总资产周转率(次)',col,safe_div(rev,avg_ta) if rev and avg_ta else None)
        s('流动资产周转率(次)',col,safe_div(rev,avg_ca) if rev and avg_ca else None)
        s('固定资产周转率(次)',col,safe_div(rev,fa) if rev and fa else None)

        # 发展
        s('营收同比(%)',col,(rev-prev_rev)/abs(prev_rev)*100 if rev and prev_rev and prev_rev!=0 else None)
        s('净利润同比(%)',col,(np_-prev_np)/abs(prev_np)*100 if np_ and prev_np and prev_np!=0 else None)
        s('总资产增长率(%)',col,(ta-prev_ta)/abs(prev_ta)*100 if ta and prev_ta and prev_ta!=0 else None)

        # 现金流
        s('经营现金流/净利润',col,safe_div(op_cf,np_) if op_cf and np_ else None)
        fc = op_cf-capex if op_cf and capex else None
        s('自由现金流(元)',col,fc)
        s('现金收入比',col,safe_div(cfs,rev) if cfs and rev else None)

        # 现金流组合
        combo = '/'.join([('+' if v>0 else '−') if v is not None else '?' for v in [op_cf,invest_cf,finance_cf]])
        s('现金流组合',col,combo)

        # Z-score
        x1 = safe_div(ca-cl,ta) if ca and cl and ta else None
        surplus = get_val(df_bs,'盈余公积',col) or 0
        retained = get_val(df_bs,'未分配利润',col) or 0
        x2 = safe_div(surplus+retained,ta) if ta else None
        ie = get_val(df_is,'其中：利息费用',col) or fin or 0
        x3 = safe_div(pt+abs(ie),ta) if pt and ta else None
        x4 = safe_div(eq,tl) if eq and tl else None
        x5 = safe_div(rev,ta) if rev and ta else None
        if all([x1,x2,x3,x4,x5]):
            z = 1.2*x1+1.4*x2+3.3*x3+0.6*x4+1.0*x5
            s('Altman Z-score',col,z)
            s('Z-score判定',col,'安全区' if z>2.99 else ('灰色区' if z>=1.81 else '危险区'))
    return pd.DataFrame(result)


# ======================== 模块2: 销量数据 ========================
def collect_sales_data():
    """收集中国汽车行业多维销量数据 - 多Sheet"""
    logger.info(f'\n{"="*60}')
    logger.info(f'  模块2: 中国汽车行业多维销量数据')
    logger.info(f'{"="*60}')

    sheets = {}

    # 1. 整体市场 (销量默认 + 产量)
    for indicator, sheet_n in [('销量','整体市场_销量'), ('产量','整体市场_产量')]:
        try:
            df = ak.car_market_total_cpca(symbol='狭义乘用车', indicator=indicator)
            if df is not None and len(df) > 0:
                sheets[sheet_n] = df
        except Exception: pass

    # 2. 燃料类型 × 4个细分 (整体/轿车/SUV/MPV)
    for symbol in ['整体市场','轿车','SUV','MPV']:
        try:
            df = ak.car_market_fuel_cpca(symbol=symbol)
            if df is not None and len(df) > 0:
                sheets[f'燃料类型_{symbol}'] = df
        except Exception: pass

    # 3. 细分市场级别 × 3 (轿车/SUV/MPV → A00/A0/A/B/C)
    for symbol in ['轿车','SUV','MPV']:
        try:
            df = ak.car_market_segment_cpca(symbol=symbol)
            if df is not None and len(df) > 0:
                sheets[f'细分市场_{symbol}'] = df
        except Exception: pass

    # 4. 厂商月度排行
    try:
        df = ak.car_market_man_rank_cpca()
        if df is not None and len(df) > 0: sheets['厂商月度排行'] = df
    except Exception: pass

    # 5. 国别品牌销量 (自主/德系/日系/美系/韩系/法系) + 自主占比计算
    try:
        df = ak.car_market_country_cpca()
        if df is not None and len(df) > 0:
            total_cols = [c for c in df.columns if c != '月份']
            df['自主占比(%)'] = df.apply(
                lambda r: round(r['自主']/sum(r[c] for c in total_cols)*100,1), axis=1)
            sheets['国别品牌销量'] = df
    except Exception: pass

    # 6. 盖世汽车品牌销量排行 (50品牌)
    try:
        df = ak.car_sale_rank_gasgoo()
        if df is not None and len(df) > 0: sheets['盖世品牌排行'] = df
    except Exception: pass
    if sheets:
        path = DATA_DIR / '中国汽车行业多维销量数据库.xlsx'
        writer = pd.ExcelWriter(path, engine='openpyxl')
        for sn, df in sheets.items():
            df.to_excel(writer, sheet_name=sn[:31], index=False)
        writer.close()
        total_rows = sum(df.shape[0] for df in sheets.values())
        logger.info(f'  ✅ {path.name}: {len(sheets)} Sheets, {total_rows}行 ({os.path.getsize(path)/1024:.0f}KB)')
        return path
    return None


# ======================== Excel 写入 ========================

def write_statement_sheet(df, sheet_name, writer):
    if df is None or df.empty: return
    df.to_excel(writer, sheet_name=sheet_name[:31], startrow=0)
    ws = writer.sheets[sheet_name[:31]]

    hf = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    hfont = Font(bold=True, size=10, color='FFFFFF')
    sf = PatternFill(start_color='D9E2F3', end_color='D9E2F3', fill_type='solid')
    sfont = Font(bold=True, size=9, color='1F3864')
    wfont = Font(bold=True, size=9, color='C00000')
    b = Border(left=Side('thin','D9D9D9'), right=Side('thin','D9D9D9'),
               top=Side('thin','D9D9D9'), bottom=Side('thin','D9D9D9'))

    for col in range(1, ws.max_column + 1):
        c = ws.cell(row=1, column=col)
        c.font, c.fill, c.alignment = hfont, hf, Alignment(horizontal='center', vertical='center')
        c.border = b

    for row in range(2, ws.max_row + 1):
        ca = ws.cell(row=row, column=1)
        val = str(ca.value) if ca.value else ''
        if val.startswith(('一、','二、','三、','四、','五、','六、','七、')):
            ca.font, ca.fill = sfont, sf
        elif any(k in val for k in ['危险','预警']):
            ca.font = wfont
        else: ca.font = Font(size=9)
        ca.alignment = Alignment(horizontal='left', vertical='center')
        ca.border = b
        for col in range(2, ws.max_column + 1):
            c = ws.cell(row=row, column=col)
            c.alignment = Alignment(horizontal='right', vertical='center')
            c.border = b
            if c.value is not None and isinstance(c.value, (int, float)):
                rn = str(ws.cell(row=row, column=1).value) or ''
                if rn in EXCLUDE_FROM_YUAN: c.number_format = '0.0000'
                elif any(k in rn for k in ['率','比']): c.number_format = '0.00%'
                elif any(k in rn for k in ['次','天']): c.number_format = '0.00'
                elif abs(c.value) >= 1: c.number_format = '#,##0'
                else: c.number_format = '0.0000'
    ws.column_dimensions['A'].width = 50
    for col in range(2, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = 16
    ws.freeze_panes = 'B2'


# ======================== 单公司处理 ========================

def process_company(symbol, name):
    cf_items, bs_items, is_items = fetch_and_pivot(symbol, name)
    if not cf_items and not bs_items and not is_items:
        logger.warning(f'    ⚠️ 无数据')
        return None

    df_cf = items_to_df(cf_items, CF_ORDER)
    df_bs = items_to_df(bs_items, BS_ORDER)
    df_is = items_to_df(is_items, IS_ORDER)
    df_derived = calc_derived_metrics(df_bs, df_is, df_cf)
    logger.info(f'    CF{df_cf.shape} BS{df_bs.shape} IS{df_is.shape} D{df_derived.shape}')

    issues = []
    annual_cols = [c for c in df_bs.columns if '年报' in str(c)]
    for col in annual_cols:
        ta = get_val(df_bs,'资产合计',col) or get_val(df_bs,'*资产合计',col)
        tl = get_val(df_bs,'负债合计',col) or get_val(df_bs,'*负债合计',col)
        eq = get_val(df_bs,'所有者权益（或股东权益）合计',col) or get_val(df_bs,'*所有者权益（或股东权益）合计',col)
        if ta and tl and eq and abs(ta-(tl+eq))/ta>0.01:
            issues.append(f'[{col}] 差={abs(ta-tl-eq)/1e8:.1f}亿')

    output_path = DATA_DIR / f'{symbol}_{name}_财务数据.xlsx'
    writer = pd.ExcelWriter(output_path, engine='openpyxl')
    if not df_bs.empty: write_statement_sheet(df_bs, '资产负债表', writer)
    if not df_is.empty: write_statement_sheet(df_is, '利润表', writer)
    if not df_cf.empty: write_statement_sheet(df_cf, '现金流量表', writer)
    if not df_derived.empty: write_statement_sheet(df_derived, '财务分析指标', writer)
    df_issues = pd.DataFrame({'校验问题': issues}) if issues else pd.DataFrame({'校验结果': ['✅ 通过']})
    df_issues.to_excel(writer, sheet_name='数据质量校验', index=False)
    writer.close()
    logger.info(f'    ✅ {output_path.name} ({os.path.getsize(output_path)/1024:.0f}KB)')
    return output_path, issues


# ======================== 综合对比 ========================

def generate_comparison(results):
    logger.info(f'\n{"="*60}')
    logger.info(f'  综合对比表 (8家)')
    logger.info(f'{"="*60}')

    comp_rows = []
    for symbol, name, note, path in results:
        xl = pd.ExcelFile(path)
        if '利润表' not in xl.sheet_names: continue
        df_is = pd.read_excel(xl, '利润表', index_col=0)
        annual_cols = [c for c in df_is.columns if '年报' in str(c)]
        if not annual_cols: continue
        latest = annual_cols[-1]
        row = {'代码':symbol,'公司':name,'分类':note,'最新年报':latest}
        for item in ['营业收入','净利润','归属于母公司所有者的净利润',
                    '营业利润','利润总额','基本每股收益','稀释每股收益',
                    '营业成本','销售费用','管理费用','研发费用','财务费用']:
            if item in df_is.index:
                v = df_is.loc[item,latest]
                if pd.notna(v): row[item] = v
        if '财务分析指标' in xl.sheet_names:
            df_d = pd.read_excel(xl,'财务分析指标',index_col=0)
            for m in ['毛利率(%)','净利率(%)','ROE(%)','资产负债率(%)',
                      '经营现金流/净利润','Altman Z-score']:
                if m in df_d.index and latest in df_d.columns:
                    v = df_d.loc[m,latest]
                    if pd.notna(v): row[m] = v
        comp_rows.append(row)

    if comp_rows:
        df_comp = pd.DataFrame(comp_rows)
        if '营业收入' in df_comp.columns:
            df_comp = df_comp.sort_values('营业收入', ascending=False).reset_index(drop=True)
        df_comp.to_excel(DATA_DIR / '8家车企综合对比.xlsx', index=False)
        logger.info(f'  ✅ 8家车企综合对比.xlsx')

        key_cols = ['代码','公司','最新年报','营业收入','净利润','毛利率(%)','净利率(%)','ROE(%)']
        avail = [c for c in key_cols if c in df_comp.columns]
        logger.info(f'\n{"="*60}')
        logger.info(f'  最新年报速览 (按营收排序)')
        logger.info(f'{"="*60}')
        df_show = df_comp[avail].copy()
        for c in df_show.columns:
            if c in ('代码','公司','最新年报'): continue
            df_show[c] = df_show[c].apply(lambda x: f'{x/1e8:.1f}亿' if pd.notna(x) and abs(x)>1e6 else (f'{x:.2f}%' if pd.notna(x) and '%' in c else (f'{x:.4f}' if pd.notna(x) else '-')))
        print(df_show.to_string(index=False))


# ======================== 主流程 ========================

def main():
    logger.info(f'{"="*80}')
    logger.info(f'  财务大数据智能决策竞赛 — 统一数据采集')
    logger.info(f'  8家A股车企 | 2016Q1—2025Q4 | 会计准则: A股统一')
    logger.info(f'  输出: {DATA_DIR}')
    logger.info(f'{"="*80}\n')

    # === 模块1: 三大报表 ===
    logger.info(f'{"="*60}')
    logger.info(f'  模块1: 8家A股车企 三大报表')
    logger.info(f'{"="*60}')

    results, all_issues = [], {}
    for symbol, name, note in COMPANIES:
        result = process_company(symbol, name)
        if result:
            path, issues = result
            results.append((symbol, name, note, path))
            all_issues[name] = issues

    generate_comparison(results)

    # === 模块2: 销量数据 ===
    collect_sales_data()

    # === 数据质量 ===
    logger.info(f'\n{"="*60}')
    logger.info(f'  数据质量报告')
    logger.info(f'{"="*60}')
    for n, issues in all_issues.items():
        logger.info(f'  {n}: {"✅" if not issues else f"⚠️ {len(issues)}"}')
        for i in issues[:1]: logger.info(f'    {i}')

    logger.info(f'\n{"="*60}')
    logger.info(f'  ✅ 全部完成! 共 {len(results)} 家 文件位于 {DATA_DIR}')
    logger.info(f'{"="*60}')


if __name__ == '__main__':
    main()
